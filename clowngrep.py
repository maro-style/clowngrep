#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""ClownGrep - local domain matching for text/data-leak datasets."""

import csv
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import warnings
from collections import OrderedDict
from pathlib import Path
from typing import Iterator, Optional

SUPPORTED_EXTS = {"txt", "log", "csv", "sql", "xlsx", "xlsm", "json"}
TEXT_BLOCK_EXTS = {"txt", "log"}
DOMAIN_TOKEN_RE = re.compile(r"(?i)(?:[a-z0-9-]+\.)+[a-z0-9-]{2,63}")
MAX_PATH_COMPONENT = 100
MAX_OPEN_OUTPUT_FILES = 64
__version__ = "1.1.0"
WINDOWS_RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


def print_banner():
    print(r"""
__________________________________________________________________________________________
  ______   __                                     ______
/      \ |  \                                   /      \
|  $$$$$$\| $$  ______   __   __   __  _______  |  $$$$$$\  ______    ______    ______
| $$   \$$| $$ /      \ |  \ |  \ |  \|       \ | $$ __\$$ /      \  /      \  /      \
| $$      | $$|  $$$$$$\| $$ | $$ | $$| $$$$$$$\| $$|    \|  $$$$$$\|  $$$$$$\|  $$$$$$\
| $$   __ | $$| $$  | $$| $$ | $$ | $$| $$  | $$| $$ \$$$$| $$   \$$| $$    $$| $$  | $$
| $$__/  \| $$| $$__/ $$| $$_/ $$_/ $$| $$  | $$| $$__| $$| $$      | $$$$$$$$| $$__/ $$
\$$    $$| $$ \$$    $$ \$$   $$   $$| $$  | $$ \$$    $$| $$       \$$     \| $$    $$
  \$$$$$$  \$$  \$$$$$$   \$$$$$\$$$$  \$$   \$$  \$$$$$$  \$$        \$$$$$$$| $$$$$$$
                                                                              | $$
                                        |0                                    | $$
                                         \`.     ___                          | $$
                                          \ \   / __>0                        \$$
                                      /\  /  |/' /
                                    /  \/   `  ,`'--.
                                  / /(___________)_ \
                                  |/ //.-.   .-.\\ \ \
                                  0 // :@ ___ @: \\ \/
                                    ( o ^(___)^ o ) 0
                                     \ \_______/ /
                                 /\   '._______.'--.
                                 \ /|  |<_____>    |
                                  \ \__|<_____>____/|__
                                   \____<_____>_______/
                                       |<_____>    |
                                       |<_____>    |
                                       :<_____>____:
                                      / <_____>   /|
                                     /  <_____>  / |
                                    /___________/  |
                                    |           | _|__
                                    |           | ---||_
                                    |  |||||||  |  | [__]
                                    |  by maro  |  /
                                    |  |||||||  | /
                                    |___________|/
__________________________________________________________________________________________
""")

def clean_name(name: str) -> str:
    """Return a filesystem-safe path component on Unix and Windows."""
    cleaned = re.sub(r"[^A-Za-z0-9_.\-]", "_", str(name).strip())
    cleaned = cleaned.strip(" .")
    if not cleaned or cleaned in {".", ".."}:
        cleaned = "_"
    if cleaned.upper() in WINDOWS_RESERVED_NAMES:
        cleaned += "_"
    return cleaned[:MAX_PATH_COMPONENT]


def input_key(file_path: Path, input_folder: Path) -> str:
    """Create a stable output key while retaining the source subdirectory."""
    try:
        rel = file_path.relative_to(input_folder).with_suffix("")
    except ValueError:
        rel = Path(file_path.stem)

    safe_parts = [clean_name(part) for part in rel.parts]
    key = "__".join(safe_parts) or "_"
    if len(key) <= 180:
        return key

    digest = hashlib.sha256(rel.as_posix().encode("utf-8", errors="ignore")).hexdigest()[:10]
    return f"{key[:165]}__{digest}"


def human_size(num_bytes: int) -> str:
    units = ["B", "KB", "MB", "GB", "TB"]
    n = float(num_bytes)
    for unit in units:
        if n < 1024 or unit == units[-1]:
            return f"{n:.1f}{unit}"
        n /= 1024.0
    return f"{n:.1f}TB"


def render_bar(current: int, total: int, width: int = 28) -> str:
    total = max(total, 1)
    ratio = min(max(current / total, 0.0), 1.0)
    filled = int(width * ratio)
    return "█" * filled + "-" * (width - filled)


def detect_encoding(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            with path.open("r", encoding=enc, errors="strict") as f:
                f.read(8192)
            return enc
        except (UnicodeDecodeError, OSError):
            continue
    return "latin-1"


def sniff_dialect(path: Path, enc: str):
    with path.open("r", encoding=enc, errors="ignore", newline="") as f:
        sample = f.read(8192)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def is_header_row(cells) -> bool:
    joined = " | ".join((c or "").strip().lower() for c in cells[:3])
    markers = (
        "third parties", "third party", "customer", "value",
        "terza parte", "terze parti", "cliente", "dominio",
    )
    return any(marker in joined for marker in markers)


def normalize_domain(value: str) -> Optional[str]:
    domain = (value or "").strip().casefold().lstrip("@").strip(".")
    if not domain or len(domain) > 253 or "." not in domain:
        return None

    labels = domain.split(".")
    for label in labels:
        if not 1 <= len(label) <= 63:
            return None
        if label.startswith("-") or label.endswith("-"):
            return None
        if not re.fullmatch(r"[a-z0-9-]+", label):
            return None

    if len(labels[-1]) < 2:
        return None
    return domain


def load_clienti_csv(csv_path: Path):
    enc = detect_encoding(csv_path)
    dialect = sniff_dialect(csv_path, enc)

    domain_map = {}
    domains = set()

    with csv_path.open("r", encoding=enc, errors="ignore", newline="") as f:
        reader = csv.reader(f, dialect=dialect)
        first_row_checked = False

        for row_no, row in enumerate(reader, start=1):
            if not row:
                continue

            row = [("" if c is None else str(c)).strip() for c in row]
            if len(row) < 3:
                continue

            if not first_row_checked:
                first_row_checked = True
                if is_header_row(row):
                    continue

            third_of = row[0].strip()
            customer = row[1].strip()
            raw_value = row[2]
            value = normalize_domain(raw_value)

            if not customer and not raw_value:
                continue
            if not customer:
                raise RuntimeError(f"Riga {row_no}: nome cliente/terza parte mancante")
            if not value:
                raise RuntimeError(f"Riga {row_no}: dominio non valido: {raw_value!r}")

            if third_of:
                owner = ("third_party", third_of, customer)
            else:
                owner = ("client", customer, None)

            owners = domain_map.setdefault(value, set())
            if owner in owners:
                # Riga duplicata identica: non aggiunge una seconda relazione.
                domains.add(value)
                continue

            # Modello molti-a-molti: ogni riga valida rappresenta una relazione
            # indipendente. Lo stesso dominio puo' quindi essere associato a:
            # - piu' clienti diretti;
            # - piu' terze parti;
            # - una combinazione di clienti diretti e terze parti.
            # Il set evita solamente di duplicare una relazione identica.
            owners.add(owner)
            domains.add(value)

    if not domain_map:
        raise RuntimeError("Nessun dominio valido trovato in clienti.csv")

    return domain_map, sorted(domains)


def _owner_set(value):
    """Normalize a domain-map value to a set of owner tuples.

    ClownGrep 1.1 uses sets because a domain may map to multiple
    third-party relationships and each third party may belong to multiple customers. The tuple fallback keeps compatibility with mappings
    created by older code or external callers.
    """
    if value is None:
        return set()
    if isinstance(value, tuple) and len(value) == 3 and value[0] in {"client", "third_party"}:
        return {value}
    return set(value)


def host_to_owners(host: str, domain_map: dict):
    """Return all owner relationships matching a host or one of its parent domains."""
    host = (host or "").casefold().strip().strip(".")
    if not host:
        return set()

    if host in domain_map:
        return _owner_set(domain_map[host])

    parts = host.split(".")
    for i in range(1, len(parts) - 1):
        candidate = ".".join(parts[i:])
        if candidate in domain_map:
            return _owner_set(domain_map[candidate])

    return set()


def host_to_owner(host: str, domain_map: dict):
    """Backward-compatible helper for domains with exactly one owner.

    Returns the single owner when unambiguous, otherwise None. New code should
    use host_to_owners() because third-party domains may map to multiple clients.
    """
    owners = host_to_owners(host, domain_map)
    if len(owners) == 1:
        return next(iter(owners))
    return None


class OutputHandlePool:
    """LRU pool that prevents one open file descriptor per matched customer."""

    def __init__(self, output_folder: Path, max_open: int = MAX_OPEN_OUTPUT_FILES):
        self.output_folder = output_folder
        self.max_open = max(1, max_open)
        self._handles = OrderedDict()
        self._targets = set()

    @property
    def target_count(self) -> int:
        return len(self._targets)

    def _path_for(self, owner_kind: str, client_name: str, third_party_name: Optional[str], input_stem: str) -> Path:
        root_clienti = self.output_folder / "clienti"
        root_terze = self.output_folder / "terze_parti"
        safe_client = clean_name(client_name)

        if owner_kind == "third_party" and third_party_name:
            out_dir = root_terze / safe_client / clean_name(third_party_name)
        else:
            out_dir = root_clienti / safe_client

        out_dir.mkdir(parents=True, exist_ok=True)
        return out_dir / f"{input_stem}.txt"

    def write(self, owner, input_stem: str, text: str) -> None:
        owner_kind, client_name, third_party_name = owner
        key = (owner_kind or "", client_name or "", third_party_name or "", input_stem)
        self._targets.add(key)

        handle = self._handles.pop(key, None)
        if handle is None:
            out_path = self._path_for(owner_kind, client_name, third_party_name, input_stem)
            handle = out_path.open("a", encoding="utf-8", errors="ignore")
        self._handles[key] = handle

        while len(self._handles) > self.max_open:
            _, old_handle = self._handles.popitem(last=False)
            old_handle.close()

        handle.write(text)

    def close(self) -> None:
        while self._handles:
            _, handle = self._handles.popitem(last=False)
            try:
                handle.close()
            except OSError:
                pass


class StreamingReport:
    """Write report sections incrementally to avoid retaining all matches in RAM."""

    def __init__(self, report_path: Path):
        self.report_path = report_path
        report_path.parent.mkdir(parents=True, exist_ok=True)
        self._client_tmp = tempfile.NamedTemporaryFile(
            mode="w", delete=False, dir=report_path.parent,
            prefix=".clowngrep_clients_", suffix=".tmp", encoding="utf-8"
        )
        self._third_tmp = tempfile.NamedTemporaryFile(
            mode="w", delete=False, dir=report_path.parent,
            prefix=".clowngrep_third_", suffix=".tmp", encoding="utf-8"
        )
        self._client_count = 0
        self._third_count = 0
        self._closed = False

    def add(self, owner_kind, client_name, third_party_name, source_file, domain_found, line_text) -> None:
        if owner_kind == "client":
            target = self._client_tmp
            self._client_count += 1
            target.write("------------------------------------------------------------\n")
            target.write(f"cliente        : {client_name}\n")
            target.write(f"file sorgente  : {source_file}\n")
            target.write(f"dominio        : {domain_found}\n")
            target.write(f"testo trovato  : {line_text.strip()}\n\n")
        else:
            target = self._third_tmp
            self._third_count += 1
            target.write("------------------------------------------------------------\n")
            target.write(f"cliente        : {client_name}\n")
            target.write(f"terza parte    : {third_party_name or '-'}\n")
            target.write(f"file sorgente  : {source_file}\n")
            target.write(f"dominio        : {domain_found}\n")
            target.write(f"testo trovato  : {line_text.strip()}\n\n")

    @staticmethod
    def _copy_text(src_path: Path, dst_handle) -> None:
        with src_path.open("r", encoding="utf-8", errors="ignore") as src:
            shutil.copyfileobj(src, dst_handle)

    def finalize(self) -> None:
        if self._closed:
            return
        self._client_tmp.close()
        self._third_tmp.close()

        client_path = Path(self._client_tmp.name)
        third_path = Path(self._third_tmp.name)
        try:
            with self.report_path.open("w", encoding="utf-8", errors="ignore") as f:
                f.write("OUTPUT\n")
                f.write("========================================\n\n")
                f.write("CLIENTI\n")
                f.write("========================================\n\n")
                if self._client_count:
                    self._copy_text(client_path, f)
                else:
                    f.write("Nessun match clienti.\n\n")

                f.write("TERZE PARTI\n")
                f.write("========================================\n\n")
                if self._third_count:
                    self._copy_text(third_path, f)
                else:
                    f.write("Nessun match terze parti.\n")
        finally:
            client_path.unlink(missing_ok=True)
            third_path.unlink(missing_ok=True)
            self._closed = True

    def abort(self) -> None:
        if self._closed:
            return
        for handle in (self._client_tmp, self._third_tmp):
            try:
                handle.close()
            except OSError:
                pass
            Path(handle.name).unlink(missing_ok=True)
        self._closed = True


def build_patterns_file(domains) -> Path:
    """Create a private temporary pattern file and return its path."""
    tmp = tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", prefix="clowngrep_patterns_", delete=False, encoding="utf-8"
    )
    try:
        for domain in domains:
            tmp.write(domain + "\n")
        tmp.flush()
    finally:
        tmp.close()

    path = Path(tmp.name)
    try:
        path.chmod(0o600)
    except OSError:
        pass
    return path


def pick_search_engine() -> str:
    if shutil.which("rg"):
        return "rg"
    if shutil.which("grep"):
        return "grep"
    raise RuntimeError("Nessun motore di ricerca disponibile: installa ripgrep (rg) o grep.")


def get_managed_venv_python(venv_dir: Path) -> Path:
    """Return the Python interpreter path for ClownGrep's managed virtualenv."""
    if os.name == "nt":
        return venv_dir / "Scripts" / "python.exe"
    return venv_dir / "bin" / "python"


def running_inside_venv() -> bool:
    return getattr(sys, "base_prefix", sys.prefix) != sys.prefix


def install_excel_dependency(python_executable: Path, base: Path) -> None:
    """Install project dependencies using the supplied Python interpreter."""
    requirements = base / "requirements.txt"
    if requirements.is_file():
        cmd = [str(python_executable), "-m", "pip", "install", "-r", str(requirements)]
    else:
        cmd = [str(python_executable), "-m", "pip", "install", "openpyxl"]

    print("[INFO] Installo automaticamente le dipendenze Excel...")
    result = subprocess.run(cmd, text=True)
    if result.returncode != 0:
        raise RuntimeError(
            "Installazione automatica di openpyxl fallita. "
            "Controlla la connessione Internet e i messaggi di pip mostrati sopra."
        )


def ensure_excel_dependency(base: Path) -> None:
    """Ensure openpyxl is available, creating a local venv when needed.

    On Homebrew/PEP 668 Python installations this avoids modifying the
    externally managed system interpreter. If ClownGrep is already running
    inside a virtual environment, the dependency is installed in that
    environment instead. Otherwise a private .clowngrep_venv is created next
    to the script and ClownGrep restarts itself with that interpreter.
    """
    try:
        import openpyxl  # noqa: F401
        return
    except ImportError:
        pass

    if running_inside_venv():
        install_excel_dependency(Path(sys.executable), base)
        try:
            import importlib
            importlib.invalidate_caches()
            import openpyxl  # noqa: F401
            return
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl risulta installato, ma non e' importabile nell'ambiente virtuale corrente."
            ) from exc

    venv_dir = base / ".clowngrep_venv"
    venv_python = get_managed_venv_python(venv_dir)

    if not venv_python.exists():
        print(f"[INFO] openpyxl non trovato. Creo ambiente virtuale locale: {venv_dir.name}")
        result = subprocess.run([sys.executable, "-m", "venv", str(venv_dir)], text=True)
        if result.returncode != 0:
            raise RuntimeError(
                "Creazione dell'ambiente virtuale fallita. "
                "Verifica che il modulo Python 'venv' sia disponibile."
            )

    check = subprocess.run(
        [str(venv_python), "-c", "import openpyxl"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    if check.returncode != 0:
        install_excel_dependency(venv_python, base)

    check = subprocess.run(
        [str(venv_python), "-c", "import openpyxl"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    if check.returncode != 0:
        raise RuntimeError("openpyxl non e' disponibile nella virtualenv locale dopo l'installazione.")

    print("[INFO] Dipendenze Excel pronte. Riavvio ClownGrep nell'ambiente locale...\n")
    os.execv(str(venv_python), [str(venv_python), str(Path(__file__).resolve()), *sys.argv[1:]])


def search_file(engine: str, file_path: Path, patterns_file: Path) -> Iterator[str]:
    """Stream matching lines from rg/grep instead of buffering all results in RAM."""
    env = dict(os.environ)
    env["LC_ALL"] = "C"

    if engine == "rg":
        cmd = [
            "rg", "--no-heading", "--no-filename", "--line-number",
            "-i", "-F", "-f", str(patterns_file), str(file_path),
        ]
    else:
        cmd = [
            "grep", "-a", "-h", "-n", "-i", "-F", "-f",
            str(patterns_file), str(file_path),
        ]

    proc = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        errors="ignore",
        env=env,
    )
    assert proc.stdout is not None
    completed = False
    try:
        for line in proc.stdout:
            yield line.rstrip("\n")
        completed = True
    finally:
        proc.stdout.close()
        if not completed and proc.poll() is None:
            proc.terminate()
            try:
                proc.wait(timeout=2)
            except subprocess.TimeoutExpired:
                proc.kill()
                proc.wait()

    stderr = ""
    if proc.stderr is not None:
        stderr = proc.stderr.read().strip()
        proc.stderr.close()
    returncode = proc.wait()
    if returncode not in (0, 1):
        raise RuntimeError(stderr or f"Errore eseguendo {engine} su {file_path}")


def format_elapsed(seconds: float) -> str:
    seconds = max(0, int(seconds))
    hours, rem = divmod(seconds, 3600)
    minutes, secs = divmod(rem, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def run_with_heartbeat(label: str, func, interval: float = 2.0):
    """Run a blocking function while periodically showing liveness."""
    stop_event = threading.Event()
    started = time.monotonic()

    def heartbeat():
        while not stop_event.wait(interval):
            elapsed = format_elapsed(time.monotonic() - started)
            print(f"\r   {label} | in lavorazione | tempo {elapsed}", end="", flush=True)

    worker = threading.Thread(target=heartbeat, daemon=True)
    print(f"   {label} ...", flush=True)
    worker.start()
    ok = False
    try:
        result = func()
        ok = True
        return result
    finally:
        stop_event.set()
        worker.join(timeout=interval + 0.5)
        elapsed = format_elapsed(time.monotonic() - started)
        status = "completato" if ok else "ERRORE"
        print(f"\r   {label} | {status} | tempo {elapsed}" + " " * 24, flush=True)


def xlsx_to_temp_text(file_path: Path) -> Path:
    import openpyxl

    tmp = tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False, encoding="utf-8")
    wb = None
    started = time.monotonic()

    try:
        def open_workbook():
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                return openpyxl.load_workbook(file_path, data_only=True, read_only=True)

        wb = run_with_heartbeat("[EXCEL 1/3] Apertura workbook", open_workbook)

        sheets = list(wb.worksheets)
        total_sheets = len(sheets)
        sheet_rows = [max(0, int(sheet.max_row or 0)) for sheet in sheets]
        total_rows = sum(sheet_rows)
        processed_rows = 0
        last_update = 0.0

        print(
            f"   [EXCEL 2/3] Conversione in testo: {total_sheets} fogli, "
            f"~{total_rows:,} righe dichiarate".replace(",", "."),
            flush=True,
        )

        for sheet_no, sheet in enumerate(sheets, start=1):
            expected_rows = sheet_rows[sheet_no - 1]
            sheet_started_rows = processed_rows
            progress_line_active = False

            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = ["" if cell is None else str(cell) for cell in row]
                line = f"[{sheet.title}] riga {row_idx}: " + " | ".join(cells)
                tmp.write(line + "\n")
                processed_rows += 1

                now = time.monotonic()
                if now - last_update >= 1.0:
                    elapsed = max(now - started, 0.001)
                    speed = processed_rows / elapsed
                    if total_rows > 0:
                        pct = min(100.0, processed_rows * 100.0 / total_rows)
                        filled = min(30, int(30 * processed_rows / total_rows))
                        bar = "#" * filled + "-" * (30 - filled)
                        global_txt = f"{processed_rows:,}/{total_rows:,}".replace(",", ".")
                        pct_txt = f"{pct:6.2f}%"
                    else:
                        bar = "?" * 30
                        global_txt = f"{processed_rows:,}/?".replace(",", ".")
                        pct_txt = "   n/d "

                    sheet_done = processed_rows - sheet_started_rows
                    sheet_total_txt = f"{expected_rows:,}".replace(",", ".") if expected_rows else "?"
                    temp_mb = tmp.tell() / (1024 * 1024)
                    eta_txt = (
                        format_elapsed(max(0, total_rows - processed_rows) / speed)
                        if total_rows > 0 and speed > 0 else "n/d"
                    )

                    status = (
                        f"\r   [{bar}] {pct_txt} | "
                        f"foglio {sheet_no}/{total_sheets}: {sheet.title[:28]} | "
                        f"riga {sheet_done:,}/{sheet_total_txt} | globale {global_txt} | "
                        f"{speed:,.0f} righe/s | trascorso {format_elapsed(elapsed)} | "
                        f"ETA {eta_txt} | tmp {temp_mb:.1f}MB"
                    ).replace(",", ".")
                    print(status, end="", flush=True)
                    progress_line_active = True
                    last_update = now

            # Chiude la riga di progresso prima del riepilogo del foglio.
            # In questo modo l'output resta leggibile anche quando viene copiato
            # dal terminale o reindirizzato in un file di log.
            if progress_line_active:
                print(flush=True)

            now = time.monotonic()
            elapsed = max(now - started, 0.001)
            speed = processed_rows / elapsed
            total_txt = f"{total_rows:,}".replace(",", ".") if total_rows else "?"
            pct = min(100.0, processed_rows * 100.0 / total_rows) if total_rows else 0.0
            print(
                f"\r   [EXCEL] foglio {sheet_no}/{total_sheets} completato: {sheet.title} | "
                f"globale {processed_rows:,}/{total_txt} | {pct:6.2f}% | "
                f"{speed:,.0f} righe/s | {format_elapsed(elapsed)}".replace(",", ".")
                + " " * 20,
                flush=True,
            )

        tmp.flush()
        elapsed = time.monotonic() - started
        print(
            f"   [EXCEL 2/3] Conversione completata: {processed_rows:,} righe in "
            f"{format_elapsed(elapsed)} - file temporaneo {human_size(Path(tmp.name).stat().st_size)}".replace(",", "."),
            flush=True,
        )
    except Exception:
        tmp.close()
        Path(tmp.name).unlink(missing_ok=True)
        raise
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
        tmp.close()

    return Path(tmp.name)


def extract_body_from_search_line(line: str) -> str:
    """Remove the line-number prefix emitted by rg/grep."""
    parts = line.split(":", 1)
    if len(parts) == 2 and parts[0].isdigit():
        return parts[1]
    return line


def iter_text_blocks(file_path: Path, enc: str) -> Iterator[str]:
    """Yield blank-line-delimited blocks without reading the whole file into memory."""
    block_lines = []
    with file_path.open("r", encoding=enc, errors="ignore", newline=None) as f:
        for line in f:
            line = line.rstrip("\r\n")
            if line.strip():
                block_lines.append(line)
                continue

            if block_lines:
                yield "\n".join(block_lines).strip()
                block_lines.clear()

    if block_lines:
        yield "\n".join(block_lines).strip()


def match_text(text: str, domain_map: dict):
    tokens = {token.casefold() for token in DOMAIN_TOKEN_RE.findall(text)}
    hit_owners = set()
    hit_domains = set()
    for token in tokens:
        owners = host_to_owners(token, domain_map)
        for owner in owners:
            hit_owners.add(owner)
            hit_domains.add((token, owner))
    return hit_owners, hit_domains


def process_file_text_blocks(
    file_path: Path,
    output_folder: Path,
    domain_map: dict,
    current_index: int,
    total_files: int,
    input_folder: Path,
):
    input_stem = input_key(file_path, input_folder)
    output_pool = OutputHandlePool(output_folder)
    report = StreamingReport(output_folder / f"output_{input_stem}.txt")
    matched_blocks = 0
    written_blocks = 0
    seen_hashes = {}

    bar = render_bar(current_index, total_files)
    size_txt = human_size(file_path.stat().st_size)
    display_path = file_path.relative_to(input_folder)
    print(f"[{bar}] {current_index}/{total_files} -> {display_path} ({size_txt})")

    enc = detect_encoding(file_path)
    source_file = str(display_path)
    success = False

    try:
        for block in iter_text_blocks(file_path, enc):
            owners, domains_found = match_text(block, domain_map)
            if not owners:
                continue
            matched_blocks += 1
            block_digest = hashlib.sha256(block.encode("utf-8", errors="ignore")).digest()

            for owner in owners:
                owner_hashes = seen_hashes.setdefault(owner, set())
                if block_digest in owner_hashes:
                    continue
                owner_hashes.add(block_digest)
                output_pool.write(owner, input_stem, block + "\n\n")
                written_blocks += 1

            for domain_found, owner in sorted(domains_found):
                owner_kind, client_name, third_party_name = owner
                report.add(
                    owner_kind, client_name, third_party_name,
                    source_file, domain_found, block,
                )
        success = True
    finally:
        output_pool.close()
        if success:
            report.finalize()
        else:
            report.abort()

    print(
        f"   blocchi_trovati={matched_blocks}, "
        f"blocchi_scritti={written_blocks}, "
        f"target_output={output_pool.target_count}"
    )


def process_file(
    file_path: Path,
    output_folder: Path,
    domain_map: dict,
    patterns_file: Path,
    engine: str,
    current_index: int,
    total_files: int,
    input_folder: Path,
):
    input_stem = input_key(file_path, input_folder)
    output_pool = OutputHandlePool(output_folder)
    report = StreamingReport(output_folder / f"output_{input_stem}.txt")
    matched_lines = 0
    written_lines = 0

    bar = render_bar(current_index, total_files)
    size_txt = human_size(file_path.stat().st_size)
    display_path = file_path.relative_to(input_folder)
    print(f"[{bar}] {current_index}/{total_files} -> {display_path} ({size_txt})")

    is_excel = file_path.suffix.lower().lstrip(".") in {"xlsx", "xlsm"}
    temp_path = None
    success = False

    try:
        target_path = file_path
        if is_excel:
            temp_path = xlsx_to_temp_text(file_path)
            target_path = temp_path

        line_iter = search_file(engine, target_path, patterns_file)
        if is_excel:
            # Streaming keeps memory bounded; the status line marks the start/end explicitly.
            print(f"   [EXCEL 3/3] Ricerca domini con {engine} ...", flush=True)
            search_started = time.monotonic()
        else:
            search_started = None

        for line in line_iter:
            matched_lines += 1
            rec = extract_body_from_search_line(line)
            owners, domains_found = match_text(rec, domain_map)
            if not owners:
                continue

            written_lines += 1
            for owner in owners:
                output_pool.write(owner, input_stem, rec + "\n")

            for domain_found, owner in sorted(domains_found):
                owner_kind, client_name, third_party_name = owner
                report.add(
                    owner_kind, client_name, third_party_name,
                    str(display_path), domain_found, rec,
                )

        if search_started is not None:
            print(
                f"   [EXCEL 3/3] Ricerca completata in "
                f"{format_elapsed(time.monotonic() - search_started)}",
                flush=True,
            )
        success = True
    finally:
        output_pool.close()
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)
        if success:
            report.finalize()
        else:
            report.abort()

    print(
        f"   righe_matchate={matched_lines}, "
        f"righe_scritte={written_lines}, "
        f"target_output={output_pool.target_count}"
    )


def print_mode_help(ext: Optional[str] = None) -> None:
    """Explain the processing modes with practical examples."""
    print("\n--- COME SCEGLIERE LA MODALITA' ---")
    print("1) RIGHE / DATALEAK")
    print("   ClownGrep analizza il file riga per riga.")
    print("   Se trova un dominio target, salva SOLO la riga che lo contiene.")
    print("   Usala quando ogni record e' contenuto in una singola riga.")
    print("   Esempi: CSV, SQL dump, JSON/JSONL, liste di credenziali, log lineari.")
    print("   Esempio: user=mario | host=portal.example.com | pass=***")
    print("            -> viene salvata quella singola riga.\n")
    print("2) BLOCCHI")
    print("   Disponibile solo per file .txt e .log.")
    print("   ClownGrep divide il file usando le RIGHE VUOTE come separatori.")
    print("   Se trova un dominio target in un blocco, salva TUTTO il blocco,")
    print("   incluse le righe che non contengono direttamente il dominio.")
    print("   Usala quando un singolo record occupa piu' righe.")
    print("   Esempio:")
    print("      URL: https://portal.example.com")
    print("      User: mario@example.com")
    print("      Password: ***")
    print("      Browser: Chrome")
    print("      <riga vuota>")
    print("      -> vengono salvate tutte e quattro le righe del record.\n")
    if not ext or ext in TEXT_BLOCK_EXTS:
        print("3) SCEGLI PER OGNI FILE")
        print("   Disponibile per .txt e .log.")
        print("   Non applica una sola modalita' a tutta l'estensione:")
        print("   ClownGrep ti chiedera' RIGHE o BLOCCHI per ciascun file.")
        print("   Usala quando file con la stessa estensione hanno strutture diverse.\n")
    print("REGOLA RAPIDA:")
    print("   - un record per riga                  -> scegli 1")
    print("   - un record su piu' righe, separato")
    print("     dagli altri con una riga vuota     -> scegli 2")
    if not ext or ext in TEXT_BLOCK_EXTS:
        print("   - file .txt/.log con strutture miste -> scegli 3")
    print("   - non vuoi analizzare l'estensione   -> scegli s")
    if ext and ext not in TEXT_BLOCK_EXTS:
        print(f"\nNOTA: per .{ext} e' disponibile solo la modalita' 1.")
    print("----------------------------------------\n")


def ask_mode_for_extension(ext: str, file_count: int) -> str:
    print()
    print(f"Estensione: .{ext} - file trovati: {file_count}")
    print("La scelta verra' applicata automaticamente a TUTTI i file con questa estensione.")

    if ext in TEXT_BLOCK_EXTS:
        print("1) RIGHE / dataleak -> salva solo le righe che contengono un dominio target")
        print("2) BLOCCHI          -> salva l'intero record multi-riga che contiene il dominio")
        print("                      (i blocchi devono essere separati da righe vuote)")
        print("3) SCEGLI PER FILE  -> chiedi RIGHE/BLOCCHI separatamente per ogni file")
        print("h) guida con esempi")
        print("s) salta tutti i file di questa estensione")
        prompt = "Selezione per tutti i .{} [1/2/3/h/s]: ".format(ext)
    else:
        print("1) RIGHE / dataleak -> salva solo le righe che contengono un dominio target")
        print("h) guida con esempi")
        print("s) salta tutti i file di questa estensione")
        prompt = "Selezione per tutti i .{} [1/h/s]: ".format(ext)

    while True:
        choice = input(prompt).strip().lower()
        if choice == "1":
            return "dataleak"
        if choice == "2" and ext in TEXT_BLOCK_EXTS:
            return "blocks"
        if choice == "3" and ext in TEXT_BLOCK_EXTS:
            return "per_file"
        if choice in {"h", "?", "help"}:
            print_mode_help(ext)
            continue
        if choice == "s":
            return "skip"
        if choice in {"2", "3"}:
            print("[ERRORE] Le modalita' 2 e 3 sono disponibili solo per file .txt o .log.")
        else:
            print("[ERRORE] Scelta non valida. Usa 1, 2/3 (se disponibili), h oppure s.")


def ask_mode_for_file(file_path: Path, input_folder: Path) -> str:
    """Ask how to process one TXT/LOG file when per-file mode is selected."""
    display_path = file_path.relative_to(input_folder)
    size_txt = human_size(file_path.stat().st_size)

    print()
    print(f"File: {display_path} ({size_txt})")
    print("1) RIGHE   -> ogni riga e' un record indipendente")
    print("2) BLOCCHI -> record multi-riga separati da righe vuote")
    print("   ATTENZIONE: usa BLOCCHI solo se i record sono separati da righe vuote.")
    print("   Senza separatori, l'intero file puo' essere interpretato come un unico blocco.")
    print("h) guida con esempi")
    print("s) salta solo questo file")
    prompt = "Modalita' per questo file [1/2/h/s]: "

    while True:
        choice = input(prompt).strip().lower()
        if choice == "1":
            return "dataleak"
        if choice == "2":
            return "blocks"
        if choice in {"h", "?", "help"}:
            print_mode_help(file_path.suffix.lower().lstrip("."))
            continue
        if choice == "s":
            return "skip"
        print("[ERRORE] Scelta non valida. Usa 1, 2, h oppure s.")


def build_processing_plan(files, mode_by_extension: dict, input_folder: Path):
    """Resolve extension-level choices into a concrete (file, mode) processing plan."""
    plan = []
    per_file_paths = [
        path for path in files
        if mode_by_extension[path.suffix.lower().lstrip(".")] == "per_file"
    ]

    if per_file_paths:
        print("\n=== MODALITA' PER SINGOLO FILE ===")
        print("Per i file delle estensioni impostate su 'SCEGLI PER FILE',")
        print("seleziona ora la modalita' da usare per ciascun file.")

    for path in files:
        ext = path.suffix.lower().lstrip(".")
        mode = mode_by_extension[ext]
        if mode == "skip":
            continue
        if mode == "per_file":
            mode = ask_mode_for_file(path, input_folder)
            if mode == "skip":
                continue
        plan.append((path, mode))

    return plan


def reset_output_folder(base: Path, output_folder: Path) -> None:
    """Clear only the dedicated ./outputs directory and refuse symlink escapes."""
    resolved_base = base.resolve()
    if output_folder.is_symlink():
        raise RuntimeError("La cartella outputs/ non puo' essere un link simbolico")
    resolved_output = output_folder.resolve()
    if resolved_output.parent != resolved_base or resolved_output.name != "outputs":
        raise RuntimeError(f"Percorso output non sicuro: {resolved_output}")

    if output_folder.exists():
        shutil.rmtree(output_folder)
    output_folder.mkdir(parents=True, exist_ok=True)


def main():
    base = Path(__file__).resolve().parent
    print_banner()
    print(f"ClownGrep v{__version__}\n")

    clienti_csv = base / "clienti.csv"
    input_folder = base / "file"
    output_folder = base / "outputs"

    if not clienti_csv.is_file():
        print(f"[ERRORE] Manca clienti.csv in: {base}")
        sys.exit(1)

    if not input_folder.exists():
        input_folder.mkdir(parents=True, exist_ok=True)
        print(f"[INFO] Creata cartella input: {input_folder}")

    try:
        engine = pick_search_engine()
    except RuntimeError as exc:
        print(f"[ERRORE] {exc}")
        sys.exit(1)

    try:
        domain_map, domains = load_clienti_csv(clienti_csv)
    except Exception as exc:
        print(f"[ERRORE] clienti.csv: {exc}")
        sys.exit(1)

    files = [
        path for path in sorted(input_folder.rglob("*"))
        if path.is_file()
        and not path.is_symlink()
        and path.suffix.lower().lstrip(".") in SUPPORTED_EXTS
    ]

    if not files:
        print("[ERRORE] Nessun file supportato trovato in file/ o nelle sue sottocartelle")
        sys.exit(1)

    if any(path.suffix.lower() in {".xlsx", ".xlsm"} for path in files):
        try:
            ensure_excel_dependency(base)
        except RuntimeError as exc:
            print(f"[ERRORE] {exc}")
            sys.exit(1)

    relation_count = sum(len(owners) for owners in domain_map.values())
    print(f"Motore ricerca: {engine}")
    print(f"Domini caricati: {len(domains)}")
    print(f"Relazioni dominio-owner: {relation_count}")
    print(f"File da processare: {len(files)}")
    print("\n=== INIZIO DEL CIRCO ===")

    extension_counts = {}
    for path in files:
        ext = path.suffix.lower().lstrip(".")
        extension_counts[ext] = extension_counts.get(ext, 0) + 1

    mode_by_extension = {}
    print("\n=== MODALITA' PER ESTENSIONE ===")
    for ext in sorted(extension_counts):
        mode_by_extension[ext] = ask_mode_for_extension(ext, extension_counts[ext])

    processing_plan = build_processing_plan(files, mode_by_extension, input_folder)
    skipped_files = len(files) - len(processing_plan)
    total_files = len(processing_plan)

    print("\n=== AVVIO ELABORAZIONE AUTOMATICA ===")
    print(f"File selezionati: {total_files}")
    if skipped_files:
        print(f"File esclusi: {skipped_files}")

    if not processing_plan:
        print("Nessun file selezionato. Fine.")
        return

    try:
        reset_output_folder(base, output_folder)
    except RuntimeError as exc:
        print(f"[ERRORE] {exc}")
        sys.exit(1)

    patterns_file = build_patterns_file(domains)
    try:
        for idx, (path, mode) in enumerate(processing_plan, start=1):
            try:
                if mode == "blocks":
                    process_file_text_blocks(
                        file_path=path,
                        output_folder=output_folder,
                        domain_map=domain_map,
                        current_index=idx,
                        total_files=total_files,
                        input_folder=input_folder,
                    )
                else:
                    process_file(
                        file_path=path,
                        output_folder=output_folder,
                        domain_map=domain_map,
                        patterns_file=patterns_file,
                        engine=engine,
                        current_index=idx,
                        total_files=total_files,
                        input_folder=input_folder,
                    )
            except Exception as exc:
                print(f"[ERRORE] {path.relative_to(input_folder)}: {exc}")
                print("         Continuo con il file successivo.")
    finally:
        patterns_file.unlink(missing_ok=True)

    print(f"\nOutput in: {output_folder}")


if __name__ == "__main__":
    main()
